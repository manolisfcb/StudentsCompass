from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path


OUT = Path('/Users/manuel/Documents/ENACTUS/StudentsCompass')
BLUE = '17365D'
TEAL = '0F6B78'
PALE = 'EAF2F4'
PALE2 = 'F5F7F8'
GRID = 'D9D9D9'
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_table_borders(table, color=GRID, size='6'):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = tblPr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tblPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = f'w:{edge}'
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), size)
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)


def set_col_widths(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def style_table(table, header_fill=BLUE, font_size=9.3, first_col_bold=False):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for c in table.rows[0].cells:
        set_cell_shading(c, header_fill)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(font_size)
    for r_idx, row in enumerate(table.rows[1:], start=1):
        for c_idx, c in enumerate(row.cells):
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(c)
            if r_idx % 2 == 0:
                set_cell_shading(c, PALE2)
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(font_size)
                    if first_col_bold and c_idx == 0:
                        run.font.bold = True


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0563C1')
    rPr.append(color)
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    rPr.append(underline)
    new_run.append(rPr)
    text_el = OxmlElement('w:t')
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_source(doc, label, url):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(label + ' ')
    r.bold = True
    r.font.size = Pt(9.5)
    add_hyperlink(p, url, url)
    return p


def setup_doc(title):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)
    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Aptos'
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color in [('Title', 23, BLACK), ('Heading 1', 15, BLACK), ('Heading 2', 12, BLACK), ('Heading 3', 10.5, BLACK)]:
        st = styles[name]
        st.font.name = 'Aptos Display' if name != 'Heading 3' else 'Aptos'
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(12 if name != 'Title' else 0)
        st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.keep_with_next = True
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = footer.add_run('Blue Water Bridge Pilot  |  OVIN Technology Pilot Zone  |  USD  |  16 September 2026')
    rr.font.size = Pt(8)
    rr.font.color.rgb = RGBColor(100, 100, 100)
    p = doc.add_paragraph(style='Title')
    p.add_run(title)
    p.paragraph_format.space_after = Pt(4)
    return doc


def add_meta(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    for i, (k, v) in enumerate(lines):
        r = p.add_run(f'{k}: ')
        r.bold = True
        r.font.size = Pt(10)
        r2 = p.add_run(v)
        r2.font.size = Pt(10)
        if i < len(lines) - 1:
            p.add_run('    ')


def add_intro(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(text)
    r.font.size = Pt(11)


def add_quote_summary(doc, rows):
    table = doc.add_table(rows=1, cols=4)
    hdr = ['Vendor', 'Configuration', 'Monthly Cost', '4 Month Cost']
    for i, h in enumerate(hdr):
        table.rows[0].cells[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    set_col_widths(table, [1.2, 3.55, 1.05, 1.05])
    style_table(table, header_fill=TEAL, font_size=9.2, first_col_bold=True)
    return table


def add_kv_table(doc, rows, widths=(1.55, 5.3)):
    table = doc.add_table(rows=0, cols=2)
    for key, val in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = val
    set_col_widths(table, list(widths))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for r_idx, row in enumerate(table.rows):
        for c_idx, c in enumerate(row.cells):
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(c)
            set_cell_shading(c, PALE if c_idx == 0 else (PALE2 if r_idx % 2 == 0 else 'FFFFFF'))
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(9.4)
                    run.font.color.rgb = BLACK
                    if c_idx == 0:
                        run.font.bold = True
    return table


def add_cost_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    set_col_widths(table, widths)
    style_table(table, header_fill=BLUE, font_size=8.8, first_col_bold=True)
    return table


def add_formula(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9.6)


def doc_vendor_quotes():
    doc = setup_doc('Cloud Vendor Quotes OVIN TPZ')
    add_meta(doc, [('Project', 'Blue Water Bridge Pilot'), ('Duration', '4 months'), ('Currency', 'USD'), ('Pricing date', '16 September 2026')])
    add_intro(doc, 'Este documento reúne tres configuraciones de referencia para OCI Finance. Las cifras representan consumo esperado para un piloto de cuatro meses, no un uso artificial del presupuesto aprobado. Los precios son listas públicas actuales y excluyen impuestos, descuentos negociados, soporte premium y consumos extraordinarios.')
    doc.add_heading('Resumen ejecutivo', level=1)
    doc.add_paragraph('Para la carga descrita, Supabase es la opción más simple y económica si se busca una plataforma integrada; DigitalOcean ofrece un backend y una base de datos administrada separados a un costo todavía bajo; Snowflake es el quote más costoso porque paga capacidad analítica por créditos, aunque el warehouse se suspende cuando no está en uso.')
    doc.add_paragraph('El presupuesto administrativo de aproximadamente 1,500 USD por mes debe entenderse como un techo operativo, no como consumo esperado. Ninguna de estas tres configuraciones necesita acercarse artificialmente a ese techo.')
    add_quote_summary(doc, [
        ('Supabase', 'Pro + 1 proyecto + Micro compute + 100 GB file storage incluido', '$25.00', '$100.00'),
        ('Snowflake', 'Standard, AWS Canada Central, 1 X Small, 150 h/mo, 100 GB comprimidos', '$340.00', '$1,360.00'),
        ('DigitalOcean', 'Basic Droplet 2 vCPU/4 GB + PostgreSQL 1 GiB + 100 GB Volume', '$49.15', '$196.60'),
    ])
    doc.add_paragraph('Nota de alcance: los tres quotes son alternativas, no componentes que deban sumarse entre sí.')

    doc.add_heading('Supabase Quote', level=1)
    add_kv_table(doc, [
        ('Plan', 'Pro, primer proyecto incluido'),
        ('Compute', 'Micro: 2-core ARM CPU, 1 GB RAM, $10/mo nominal; cubierto por el crédito de $10/mo del Pro Plan'),
        ('Database', 'PostgreSQL dedicado; 8 GB de disk size incluido por proyecto'),
        ('Storage', '100 GB de file storage incluido; exceso a $0.0213/GB'),
        ('Principales servicios', 'PostgreSQL, Auth, APIs ilimitadas, Storage, Realtime, Edge Functions, backups diarios 7 días, logs 7 días, métricas endpoint y email support'),
        ('Estimated Monthly Cost', '$25.00 USD'),
        ('4 Month Total', '$100.00 USD'),
    ])
    add_formula(doc, '$25 Pro + $10 Micro − $10 compute credit + $0 incremental storage dentro de 100 GB = $25.00/mes')
    doc.add_paragraph('El Micro es suficiente para un piloto con carga pequeña o moderada y evita pagar por memoria que todavía no está justificada. El Small sería una contingencia razonable si aumentan las conexiones o la concurrencia: $25 + $15 − $10 = $30/mes, pero no es necesario como configuración inicial.')
    add_source(doc, 'Fuente oficial de pricing:', 'https://supabase.com/pricing')
    add_source(doc, 'Fuente oficial de compute:', 'https://supabase.com/docs/guides/platform/manage-your-usage/compute')

    doc.add_heading('Snowflake Quote', level=1)
    add_kv_table(doc, [
        ('Edition', 'Standard'),
        ('Cloud', 'AWS'),
        ('Region', 'Canada Central'),
        ('Warehouse configuration', '1 Gen 1 X Small warehouse; auto suspend habilitado; no multi-cluster'),
        ('Estimated compute usage', '150 warehouse hours/month combinadas; rango de sensibilidad 100–200 h/month'),
        ('Warehouse size', 'X Small = 1 credit/hour; facturación por segundo con mínimo de 60 segundos'),
        ('Credits/month', '150'),
        ('USD/credit', '$2.25 en AWS Canada Central para Standard On Demand'),
        ('Storage', '100 GB comprimidos = 0.10 TB; $25/TB-month = $2.50/month'),
        ('Estimated Monthly Cost', '$340.00 USD'),
        ('4 Month Total', '$1,360.00 USD'),
    ])
    add_formula(doc, '150 h × 1 credit/h × $2.25/credit = $337.50 compute; 0.10 TB × $25/TB-month = $2.50 storage; total = $340.00/mes')
    doc.add_paragraph('Se usa un solo warehouse porque el piloto no requiere concurrencia sostenida. Con auto-suspend, el warehouse consume créditos únicamente durante cargas, transformaciones y consultas; 100 horas producirían aproximadamente $227.50/mes y 200 horas aproximadamente $452.50/mes, incluyendo el mismo supuesto de almacenamiento.')
    add_source(doc, 'Fuente oficial de créditos y storage:', 'https://www.snowflake.com/legal-files/CreditConsumptionTable.pdf')
    add_source(doc, 'Fuente oficial de warehouses:', 'https://docs.snowflake.com/en/user-guide/warehouses-overview')
    add_source(doc, 'Calculador oficial:', 'https://www.snowflake.com/en/pricing-options/calculator/')

    doc.add_heading('DigitalOcean Quote', level=1)
    add_kv_table(doc, [
        ('Droplet configuration', 'Basic shared CPU, Linux, 2 vCPU, 4 GiB RAM, 80 GiB SSD, 4,000 GiB transfer included, 24/7'),
        ('Managed Database configuration', 'PostgreSQL single node, 1 GiB RAM, 1 vCPU, 10 GiB minimum disk; no standby node'),
        ('Block Storage', '100 GiB Volume at $0.10/GiB-month'),
        ('Droplet', '$24.00/month'),
        ('Managed PostgreSQL', '$15.15/month'),
        ('Block Storage', '$10.00/month'),
        ('Estimated Monthly Cost', '$49.15 USD'),
        ('4 Month Total', '$196.60 USD'),
    ])
    add_formula(doc, '$24.00 Droplet + $15.15 PostgreSQL + 100 GiB × $0.10/GiB-month = $49.15/mes')
    doc.add_paragraph('La configuración usa un Droplet pequeño para FastAPI y los servicios de aplicación, dejando PostgreSQL en el servicio administrado. El tier de 1 GiB es apropiado para el piloto mientras el volumen de consultas y la concurrencia permanezcan modestos; se puede escalar después sin rediseñar la aplicación.')
    add_source(doc, 'Fuente oficial de Droplets:', 'https://www.digitalocean.com/pricing/droplets')
    add_source(doc, 'Fuente oficial de Managed PostgreSQL:', 'https://docs.digitalocean.com/products/databases/postgresql/details/pricing/')
    add_source(doc, 'Fuente oficial de Volumes:', 'https://docs.digitalocean.com/products/volumes/details/pricing/')

    doc.add_heading('Tabla comparativa', level=1)
    add_quote_summary(doc, [
        ('Supabase', 'Pro + Micro + 100 GB file storage incluido', '$25.00', '$100.00'),
        ('Snowflake', 'Standard + X Small 150 h/mo + 100 GB comprimidos', '$340.00', '$1,360.00'),
        ('DigitalOcean', '2 vCPU/4 GB + PostgreSQL 1 GiB + 100 GB Volume', '$49.15', '$196.60'),
    ])
    doc.add_heading('Supuestos comunes y exclusiones', level=1)
    for text in [
        'Los precios están en USD y se consultaron el 16 de septiembre de 2026. Son precios públicos; no incluyen impuestos, créditos promocionales, descuentos contractuales ni soporte premium.',
        'El piloto utiliza datasets históricos y raw de escala pequeña o moderada. Se evita usar 1 TB de almacenamiento o compute 24/7 para ML sin una necesidad demostrada.',
        'La ingestión, ETL ligero, anomaly detection, feature extraction, embeddings y modelos clásicos se consideran cargas periódicas. El procesamiento puede ejecutarse temporalmente o en el compute existente.',
        'Los quotes no incluyen costos de personal, desarrollo, licencias de KinesisIQ/Aporia, dominios, herramientas de colaboración ni transferencias extraordinarias fuera de las cuotas incluidas.',
    ]:
        doc.add_paragraph(text, style='List Bullet')
    return doc


def doc_aws():
    doc = setup_doc('AWS Compute Storage Estimate')
    add_meta(doc, [('Estimate', 'OVIN TPZ Blue Water Bridge Pilot'), ('Duration', '4 months'), ('Region', 'Canada Central ca-central-1'), ('Pricing date', '16 September 2026')])
    add_intro(doc, 'Esta estimación separada modela una arquitectura AWS pequeña y operable para el piloto. Se utilizan precios On Demand y no se emplean Reserved Instances ni Savings Plans, porque la duración prevista es de cuatro meses. La cifra resultante es consumo esperado de infraestructura, no el techo administrativo de 1,500 USD por mes.')
    doc.add_heading('Conclusión ejecutiva', level=1)
    doc.add_paragraph('La configuración recomendada cuesta aproximadamente $107.68 USD por mes y $430.72 USD por cuatro meses antes de impuestos. Incluso incluyendo un Application Load Balancer, CloudFront, logging y una base de datos administrada, el consumo esperado queda muy por debajo del presupuesto administrativo de aproximadamente $1,500 USD por mes. El presupuesto aprobado puede funcionar como ceiling operativo para crecimiento, contingencias y servicios no modelados; no es una razón para sobredimensionar la plataforma.')
    doc.add_paragraph('El AWS Pricing Calculator no pudo abrirse en esta sesión por el límite de uso del navegador de Codex. Por ello, este documento no afirma que exista un estimate persistido o un PDF exportado desde AWS; reproduce de forma auditable la metodología del calculator con precios unitarios oficiales y deja el enlace para que Finance pueda recrearlo o exportarlo directamente.')

    doc.add_heading('Arquitectura propuesta', level=1)
    doc.add_paragraph('Frontend web estático en S3 servido por CloudFront; API FastAPI y servicios de aplicación en una instancia EC2; PostgreSQL en RDS Single-AZ; datasets raw, históricos, intermedios y artefactos de modelos en S3; ingestión y ETL ligero con Lambda disparada por EventBridge; procesamiento batch y ML ejecutado primero en la EC2 existente; métricas, logs y alarmas en CloudWatch.')
    add_cost_table(doc, ['Capa', 'Servicio', 'Rol en el piloto'], [
        ('Frontend', 'S3 + CloudFront', 'Hosting estático y distribución ligera'),
        ('Backend', 'EC2 t3.medium', 'FastAPI, KinesisIQ APIs, REST y servicios de aplicación'),
        ('Database', 'RDS PostgreSQL', 'Persistencia relacional separada del backend'),
        ('Storage', 'S3 Standard', 'Raw, históricos, CSV/JSON/Parquet, intermedios y modelos'),
        ('Automation', 'EventBridge + Lambda', 'Jobs programados, ingesta y ETL ligero'),
        ('Monitoring', 'CloudWatch', 'Logs, métricas custom y alarmas'),
        ('Edge/API', 'ALB', 'Terminación HTTPS, health checks y routing del backend'),
    ], [1.15, 1.55, 4.1])

    doc.add_heading('AWS quote', level=1)
    doc.add_paragraph('Configuración de referencia para reproducir en AWS Pricing Calculator: región Canada Central, precios On Demand, Linux, 730 horas por mes. Las cantidades mensuales son supuestos de uso, no compromisos de compra.')
    rows = [
        ('Amazon EC2', '1 × t3.medium Linux', '730 h/mo × $0.0464/h', '$33.87', '$135.48'),
        ('Amazon EBS', '50 GB gp3', '50 GB × $0.088/GB-mo', '$4.40', '$17.60'),
        ('Amazon RDS', 'db.t3.micro PostgreSQL Single-AZ', '730 h/mo × $0.0200/h', '$14.60', '$58.40'),
        ('RDS storage', '50 GB General Purpose SSD', '50 GB × $0.115/GB-mo', '$5.75', '$23.00'),
        ('Amazon S3', '250 GB S3 Standard', '250 GB × $0.025/GB-mo', '$6.25', '$25.00'),
        ('Amazon CloudFront', '50 GB egress + 1M HTTPS requests', '50 × $0.085 + 1M × $0.001', '$5.25', '$21.00'),
        ('AWS Lambda', '500,000 requests; 512 MB; 1 sec avg', '250,000 GB-s × $0.0000166667 + $0.10 requests', '$4.27', '$17.08'),
        ('Amazon EventBridge', 'Scheduled rules, low event volume', 'Scheduled control-plane use; no material charge modelled', '$0.00', '$0.00'),
        ('Amazon CloudWatch', '15 GB logs + 10 metrics + 8 alarms', '10 GB billable logs × $0.50 + 5 GB archive × $0.03', '$5.15', '$20.60'),
        ('Application Load Balancer', '1 ALB + 1 average LCU', '730 h × ($0.02475 + $0.0088)', '$24.49', '$97.96'),
        ('Public IPv4', '1 in-use public IPv4 for EC2', '730 h × $0.005/h', '$3.65', '$14.60'),
        ('Batch and ML', 'Use existing EC2; no always-on ML host', 'No incremental compute in baseline', '$0.00', '$0.00'),
    ]
    add_cost_table(doc, ['Servicio', 'Configuración', 'Cálculo', 'Monthly', '4 months'], rows, [1.15, 2.0, 2.35, 0.72, 0.78])
    doc.add_paragraph('Estimated Monthly Cost: $107.68 USD', style='Heading 2')
    doc.add_paragraph('4 Month Total: $430.72 USD', style='Heading 2')
    add_formula(doc, 'Total mensual = 33.87 + 4.40 + 14.60 + 5.75 + 6.25 + 5.25 + 4.27 + 0 + 5.15 + 24.49 + 3.65 + 0 = $107.68')

    doc.add_heading('Decisiones de sizing', level=1)
    doc.add_paragraph('EC2 t3.medium es el punto de partida: 2 vCPU y 4 GiB de RAM para FastAPI, REST APIs y servicios de aplicación. No se justifica t3.large mientras no haya evidencia de presión sostenida de memoria o CPU; el salto a t3.large duplicaría aproximadamente el componente EC2 y añadiría 4 GiB de RAM sin necesidad demostrada.')
    doc.add_paragraph('RDS db.t3.micro Single-AZ con 50 GB es suficiente para un piloto con carga moderada y evita la prima de Multi-AZ. En una implementación final, este componente puede sustituirse por Supabase o Neon si se prioriza menor operación de base de datos.')
    doc.add_paragraph('No se agrega SageMaker permanente. scikit-learn, XGBoost, anomaly detection, feature extraction y embeddings se ejecutan sobre la EC2 existente o en tareas temporales. Si la medición real exige un worker separado durante 60 horas/mes, una referencia t3.medium añadiría aproximadamente 60 × $0.0464 = $2.78/mes; ese escenario no está incluido en el baseline.')
    doc.add_paragraph('El ALB se incluye en el quote porque el piloto expone APIs y requiere una entrada HTTPS estable con health checks. Si el equipo decide publicar directamente un único EC2 detrás de un reverse proxy, el costo estimado disminuiría aproximadamente $24.49/mes y el baseline quedaría en $83.19/mes.')

    doc.add_heading('Supuestos de consumo', level=1)
    for text in [
        '730 horas por mes para recursos persistentes, como exige el punto de partida del encargo.',
        'S3 Standard con 250 GB promedio almacenados. No se usa 1 TB arbitrariamente. Solicitudes PUT/GET y egress extraordinarios quedan fuera del cálculo base.',
        'CloudFront con 50 GB/mes de distribución y 1 millón de solicitudes HTTPS. Si el frontend permanece interno o tiene tráfico mínimo, este componente puede ser inferior.',
        'Lambda con 500,000 invocaciones/mes, 512 MB y 1 segundo promedio. El cálculo muestra precio bruto de uso; una cuenta elegible para Free Tier puede ver una factura incremental menor o cero.',
        'CloudWatch con 15 GB de logs ingeridos/mes, 5 GB retenidos como archivo, 10 custom metrics y 8 alarms. Se aplica el free tier de 5 GB de logs, 10 métricas y 10 alarmas al cálculo mostrado; fuera del free tier, el resultado sube modestamente.',
        'RDS backups automáticos dentro de la cuota gratuita de backup equivalente al storage provisionado; sin I/O extraordinario, snapshots adicionales, NAT Gateway, WAF, Route 53, Secrets Manager o soporte pagado.',
        'Los precios excluyen impuestos, créditos promocionales, descuentos negociados y transferencias de datos fuera de las cantidades declaradas.',
    ]:
        doc.add_paragraph(text, style='List Bullet')

    doc.add_heading('Sensibilidad y presupuesto', level=1)
    add_cost_table(doc, ['Escenario', 'Monthly Cost', '4 Month Cost', 'Interpretación'], [
        ('Baseline recomendado', '$107.68', '$430.72', 'Incluye ALB, CloudFront y monitoring conservador'),
        ('Sin ALB', '$83.19', '$332.76', 'Útil si un único EC2 sirve el backend directamente'),
        ('Con worker ML temporal 60 h/mo', '$110.46', '$441.84', 'Agrega $2.78/mo sobre el baseline'),
        ('Budget ceiling administrativo', '$1,500.00', '$6,000.00', 'Techo operativo; no es consumo esperado'),
    ], [1.7, 1.0, 1.0, 3.0])
    doc.add_paragraph('El baseline representa aproximadamente 7.2% del ceiling mensual de $1,500. Aun el escenario con worker temporal queda muy por debajo del rango original de $1,200–$2,000/mes solicitado para el calculator. La diferencia debe documentarse como capacidad presupuestaria disponible para crecimiento, contingencia y servicios futuros, no como una razón para añadir máquinas, storage, bases de datos o SageMaker sin evidencia técnica.')

    doc.add_heading('Fuentes oficiales', level=1)
    add_source(doc, 'AWS Pricing Calculator:', 'https://calculator.aws/')
    add_source(doc, 'AWS guide for EC2 estimates:', 'https://docs.aws.amazon.com/pricing-calculator/latest/userguide/ec2-estimates.html')
    add_source(doc, 'EC2 On Demand pricing:', 'https://aws.amazon.com/ec2/pricing/on-demand/')
    add_source(doc, 'Amazon EBS pricing and gp3:', 'https://aws.amazon.com/ebs/pricing/')
    add_source(doc, 'Amazon S3 pricing:', 'https://aws.amazon.com/s3/pricing/')
    add_source(doc, 'Amazon RDS PostgreSQL pricing:', 'https://aws.amazon.com/rds/postgresql/pricing/')
    add_source(doc, 'AWS Lambda pricing:', 'https://aws.amazon.com/lambda/pricing/')
    add_source(doc, 'Amazon EventBridge pricing:', 'https://aws.amazon.com/eventbridge/pricing/')
    add_source(doc, 'Amazon CloudWatch pricing:', 'https://aws.amazon.com/cloudwatch/pricing/')
    add_source(doc, 'Elastic Load Balancing pricing:', 'https://aws.amazon.com/elasticloadbalancing/pricing/')
    return doc


if __name__ == '__main__':
    d1 = doc_vendor_quotes()
    d1.save(OUT / 'Cloud_Vendor_Quotes_OVIN_TPZ.docx')
    d2 = doc_aws()
    d2.save(OUT / 'AWS_Compute_Storage_Estimate.docx')
    print('created', OUT / 'Cloud_Vendor_Quotes_OVIN_TPZ.docx')
    print('created', OUT / 'AWS_Compute_Storage_Estimate.docx')
